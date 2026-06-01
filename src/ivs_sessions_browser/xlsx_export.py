# flake8: noqa
# isort: skip_file
"""
Defined in xlsx_export.py

Filename:       xlsx_export.py
Author:         jole
Created:        20.05.2026

Description:    Holds helpers for exporting visible IVS sessions to XLSX.
"""

from __future__ import annotations

import io
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import defs as D
from .tui import Tui


COLOR_TO_HEX = {
    "black"     : "FF000000",
    "blue"      : "FF1C7ED6",
    "cyan"      : "FF0C8599",
    "green"     : "FF2B8A3E",
    "magenta"   : "FFAE3EC9",
    "red"       : "FFE03131",
    "white"     : "FFFFFFFF",
    "yellow"    : "FFB7791F",
}
INVISIBLE_ON_WHITE = {"white"}

HEADER_FILL    = PatternFill("solid", fgColor="FFD9EAF7")
HEADER_FONT    = Font(bold=True, color="FF000000")
HEADER_ALIGN   = Alignment(vertical="top", wrap_text=False)
CELL_ALIGN     = Alignment(vertical="top", wrap_text=False)

TABLE_START_ROW = 5
TABLE_START_COL = 2


def _selected_indices(pretty_print: str | list[str]) -> list[int]:
    if pretty_print == "ALL":
        return list(range(len(D.HEADERS)))

    indices = [
        D.FIELD_INDEX[D.PRETTY_PRINT_COLUMN_TO_FIELD[col_name]]
        for col_name in pretty_print
        if col_name in D.PRETTY_PRINT_COLUMN_TO_FIELD
    ]
    return indices or list(range(len(D.HEADERS)))
# ─── END OF _selected_indices() ───────────────────────────────────────────────


def _row_font_color(
    values: list[str],
    operator_bindings: dict[str, str],
    operator_colors: dict[str, str],
) -> str:
    op_label = values[D.FIELD_INDEX.get("op", 0)].strip()
    for op_key, label in operator_bindings.items():
        if label == op_label:
            color_name = operator_colors.get(op_key, "").lower()
            if color_name in INVISIBLE_ON_WHITE:
                return "FF000000"
            return COLOR_TO_HEX.get(color_name, "FF000000")
    return "FF000000"
# ─── END OF _row_font_color() ────────────────────────────────────────────────


def _formula_string(value: str) -> str:
    return value.replace('"', '""')
# ─── END OF _formula_string() ────────────────────────────────────────────────


def sessions_to_xlsx_bytes(
    rows: list[D.Row],
    pretty_print: str | list[str],
    operator_bindings: dict[str, str],
    operator_colors: dict[str, str],
) -> bytes:
    """
    Convert visible session rows to an XLSX workbook represented as bytes.
    """

    selected_indices = _selected_indices(pretty_print)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    ws.freeze_panes = ws.cell(row=TABLE_START_ROW + 1, column=TABLE_START_COL).coordinate

    for col_num, idx in enumerate(selected_indices, start=TABLE_START_COL):
        cell = ws.cell(row=TABLE_START_ROW, column=col_num, value=D.HEADERS[idx][0])
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN

    type_idx = D.FIELD_INDEX.get("type", 1)
    display_values: dict[tuple[int, int], str] = {}
    for row_num, (values, url, meta) in enumerate(rows, start=TABLE_START_ROW + 1):
        font_color = _row_font_color(values, operator_bindings, operator_colors)
        row_font = Font(color=font_color)

        for col_num, idx in enumerate(selected_indices, start=TABLE_START_COL):
            value = values[idx].strip()
            marker = Tui.type_marker(meta) if idx == type_idx else ""
            if marker:
                value = f"{value} {marker}".strip()

            display_values[(row_num, col_num)] = value
            if idx == D.FIELD_INDEX.get("code") and url:
                value = f'=HYPERLINK("{_formula_string(url)}","{_formula_string(value)}")'

            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.font = row_font
            cell.alignment = CELL_ALIGN

    filter_start = ws.cell(row=TABLE_START_ROW, column=TABLE_START_COL).coordinate
    filter_end = ws.cell(
        row=max(TABLE_START_ROW, ws.max_row),
        column=TABLE_START_COL + len(selected_indices) - 1,
    ).coordinate
    ws.auto_filter.ref = f"{filter_start}:{filter_end}"

    for col_num, idx in enumerate(selected_indices, start=TABLE_START_COL):
        width = D.HEADERS[idx][1] + 2
        max_content_width = max(
            len(display_values.get((row_num, col_num), str(ws.cell(row=row_num, column=col_num).value or "")))
            for row_num in range(TABLE_START_ROW, ws.max_row + 1)
        )
        ws.column_dimensions[get_column_letter(col_num)].width = min(
            max(width, max_content_width + 2),
            80,
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
# ─── END OF sessions_to_xlsx_bytes() ─────────────────────────────────────────


def write_sessions_xlsx(
    rows: list[D.Row],
    pretty_print: str | list[str],
    operator_bindings: dict[str, str],
    operator_colors: dict[str, str],
    output_stream: BinaryIO,
) -> None:
    """
    Write visible session rows as XLSX to the provided binary output stream.
    """

    output_stream.write(
        sessions_to_xlsx_bytes(rows, pretty_print, operator_bindings, operator_colors)
    )
# ─── END OF write_sessions_xlsx() ────────────────────────────────────────────
